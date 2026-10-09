"""
simulation_report.py - 2-3 page report of ONE completed simulation (PDF / DOCX / printable HTML).

Everything comes from the snapshot taken when the run completed (`build_snapshot`):
the same conditions, observations and results the dashboard used - nothing is
fetched again. Values are labelled OBSERVED (satellite / weather station data),
PREDICTED (ML model), DERIVED (computed indices), SCENARIO INPUT (user-chosen)
or SIMULATED (fire-spread model). Missing data is stated as missing; charts are
drawn only from recorded series.

Dependencies: reportlab (PDF), python-docx (DOCX), matplotlib (charts, optional:
without it the report states that charts are unavailable).
"""
from __future__ import annotations

import io
import math
from datetime import datetime, timezone
from typing import Dict, List, Optional

import numpy as np

TITLE = "Digital Twin Framework for Forest Fire Prediction"
DISCLAIMER = ("This simulation supports situational awareness only. It does not replace official fire monitoring, "
              "forest-department assessments or emergency instructions. Simulated spread uses assumed fuel "
              "parameters that are not calibrated against observed Bandipur fires.")
LIMITATIONS = [
    "Fire spread is a rate-of-spread cellular automaton using Canadian FBP System equations with an assumed "
    "dry-deciduous / grass fuel blend - uncalibrated for Bandipur; treat results as a what-if, not a forecast.",
    "Fuel load comes from the regional grid zone's NDVI (synthetic in demo mode) with simulated cell-to-cell "
    "variability; land cover does not restrict spread. Spotting (ember ignitions ahead of the front) is not modelled.",
    "Terrain slope uses a ~90 m DEM lattice when available, otherwise flat ground (stated below).",
    "ML risk is a model prediction for the grid zone, not an observation of fire.",
    "NASA FIRMS hotspots are satellite observations (375 m VIIRS pixels), not ground-verified fires.",
]
COMPASS = ["N", "NNE", "NE", "ENE", "E", "ESE", "SE", "SSE", "S", "SSW", "SW", "WSW", "W", "WNW", "NW", "NNW"]


def compass(deg) -> str:
    return COMPASS[int(round((float(deg) % 360) / 22.5)) % 16]


def _f(v, nd=1):
    try:
        x = float(v)
        return None if not math.isfinite(x) else round(x, nd)
    except (TypeError, ValueError):
        return None


def build_snapshot(result, cond: dict, setup: dict, mode: str, scenario: Optional[dict], live: Optional[dict],
                   twin=None, plan: Optional[dict] = None, actor: str = "") -> dict:
    """Reproducible snapshot of a completed run (JSON-serialisable)."""
    from src.dashboard.geo_fire_map import duration_label
    p = result.params
    fin = result.final
    m = result.metrics
    loc = setup.get("location") or {}
    risk_score = cond.get("risk_score")
    cat = None
    if twin is not None and risk_score is not None and np.isfinite(risk_score):
        try:
            cat = twin.alert_engine.classify(float(risk_score))
        except Exception:
            cat = None
    mc = getattr(twin, "model_choice", None)
    # weather actually used by the run
    if live and live.get("mode") == "live":
        rec = (live.get("weather") or {}).get("record") or {}
        weather = {"label": "OBSERVED (OpenWeatherMap)", "temp_c": _f(rec.get("temperature_c")),
                   "humidity_pct": _f(rec.get("humidity_pct"), 0), "wind_speed_ms": _f(rec.get("wind_speed_ms")),
                   "wind_from_deg": _f(rec.get("wind_deg"), 0), "source": "OpenWeatherMap",
                   "timestamp_utc": (live.get("weather") or {}).get("observed_utc") or
                   (live.get("weather") or {}).get("fetched_utc"),
                   "status": (live.get("weather") or {}).get("status")}
    elif scenario:
        weather = {"label": "SCENARIO INPUT (user-chosen what-if weather)", "temp_c": _f(scenario.get("temp_c")),
                   "humidity_pct": _f(scenario.get("humidity_pct"), 0), "wind_speed_ms": _f(scenario.get("wind_speed_ms")),
                   "wind_from_deg": _f(scenario.get("wind_from_deg"), 0), "source": "What-If scenario",
                   "timestamp_utc": None, "status": "scenario"}
    else:
        weather = {"label": "MODEL GRID ZONE (live or demo, see status)", "temp_c": _f(cond.get("temp_c")),
                   "humidity_pct": _f(cond.get("humidity_pct"), 0), "wind_speed_ms": _f(cond.get("wind_speed_ms")),
                   "wind_from_deg": _f(cond.get("wind_from_deg"), 0), "source": "regional twin zone "
                   + str(cond.get("zone_id")), "timestamp_utc": None,
                   "status": "demo" if getattr(twin, "offline", False) else "live/cached"}
    w_used = result.wind_schedule[0] if result.wind_schedule else (0.0, 0.0)
    # FIRMS
    firms = {"status": "not used (no LIVE observation hand-off)", "n_total": 0, "n_in_area": 0, "n_valid": 0,
             "detections": [], "fetched_utc": None, "latest_acq_utc": None}
    if live:
        fs = live.get("firms") or {}
        cls = (plan or {}).get("cls") or {}
        firms = {"status": fs.get("status"), "fetched_utc": fs.get("fetched_utc"),
                 "latest_acq_utc": fs.get("latest_acq_utc"), "n_total": int(fs.get("n", 0) or 0),
                 "n_in_area": int(cls.get("n_in_area", 0) or 0), "n_valid": int(cls.get("n_valid", 0) or 0),
                 "detections": [{k: d.get(k) for k in ("lat", "lon", "acq_utc", "date", "confidence", "conf",
                                                       "satellite", "sat", "instrument", "frp")
                                 if d.get(k) is not None}
                                for d in (live.get("detections") or [])[:25]]}
    # historical (training-season fire climatology of the zone), if the model has one
    hist = {"available": False, "note": "Historical fire records are not available for this zone in this session."}
    try:
        clim = getattr(twin, "climatology", None)
        if clim is not None and clim.table is not None and not clim.table.empty:
            import pandas as pd
            v = clim.transform(pd.DataFrame({"zone_id": [cond.get("zone_id")], "latitude": [cond.get("zone_lat")],
                                             "longitude": [cond.get("zone_lon")]}))
            hist = {"available": True, "zone_fire_rate": _f(v["zone_clim"].iloc[0], 4),
                    "neighbour_fire_rate": _f(v["nbr_clim"].iloc[0], 4),
                    "source": f"NASA FIRMS archive, model training seasons ({getattr(mc, 'climatology_file', '-')})",
                    "note": "Per-cell historical fire rate learned from the training seasons (DERIVED)."}
    except Exception:
        pass
    series = [{"minutes": x["minutes"], "burning_ha": round(x["burning"] * result.focus.cell_m ** 2 / 1e4, 3),
               "burned_ha": x["burned_ha"], "front_m": x["front_distance_m"], "ros": x["ros_m_per_min"]} for x in m]
    dom0 = result.initial_domain or result.domain
    snap = {
        "report_id": p.get("report_id"), "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "actor": actor, "title": TITLE, "mode": mode,
        "region": getattr(getattr(twin, "region", None), "name", None),
        "location": {"name": loc.get("name"), "lat": float(loc.get("lat", result.focus.lat)),
                     "lon": float(loc.get("lon", result.focus.lon))},
        "focus": {"width_m": result.focus.n_cols * result.focus.cell_m, "height_m": result.focus.n_rows * result.focus.cell_m,
                  "cell_m": result.focus.cell_m},
        "domain_initial": {"width_m": dom0.width_m, "height_m": dom0.height_m},
        "domain_final": {"width_m": result.domain.width_m, "height_m": result.domain.height_m,
                         "expansions": len(result.expansions)},
        "ignition": {"kind": "OBSERVED (NASA FIRMS)" if mode == "LIVE" else "HYPOTHETICAL (user)",
                     "points": (plan or {}).get("ign", {}).get("points") or setup.get("ignition_points") or [],
                     "placement": setup.get("placement"), "checks": p.get("ignition_checks") or []},
        "weather": weather,
        "wind_used": {"speed_ms": round(float(w_used[0]), 2), "from_deg": round(float(w_used[1]), 1)},
        "risk": {"score_pct": None if risk_score is None or not np.isfinite(risk_score) else round(100 * float(risk_score)),
                 "category": cat, "zone": cond.get("zone_id"), "zone_distance_km": cond.get("zone_distance_km"),
                 "model": getattr(mc, "label", None), "model_file": getattr(mc, "model_file", None)},
        "fwi": {k: _f(cond.get(k)) for k in ("ffmc", "dmc", "dc", "bui", "fwi")},
        "firms": firms, "historical": hist,
        "results": {"duration_min": result.duration_minutes, "duration_label": duration_label(result.duration_minutes),
                    "burned_ha": fin.get("burned_ha", 0.0), "fire_area_ha": fin.get("fire_area_ha", 0.0),
                    "burning_ha": round(fin.get("burning", 0) * result.focus.cell_m ** 2 / 1e4, 3),
                    "perimeter_m": fin.get("perimeter_m"), "front_distance_m": fin.get("front_distance_m"),
                    "peak_ros": max((x["ros_m_per_min"] for x in m), default=0.0),
                    "spread_towards": compass(w_used[1] + 180), "extinguished_at_min": p.get("extinguished_at_min"),
                    "left_focus": bool(fin.get("left_focus")), "max_kw_m": max((x.get("max_intensity_kw_m") or 0)
                                                                              for x in m) if m else 0},
        "model": {k: p.get(k) for k in ("model", "fuel_model", "ros_head_m_per_min", "ros_flank_m_per_min",
                                        "ros_back_m_per_min", "length_to_breadth", "isi", "frame_minutes")},
        "terrain": result.terrain_source,
        "series": series,
        "limitations": LIMITATIONS,
    }
    return snap


# ── charts (only from recorded series) ──────────────────────────────────────

def charts(snap: dict) -> Dict[str, bytes]:
    out: Dict[str, bytes] = {}
    s = snap.get("series") or []
    if len(s) < 2:
        return out
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return out
    t = [x["minutes"] for x in s]
    fig, ax = plt.subplots(figsize=(6.4, 2.1), dpi=150)
    ax.plot(t, [x["burned_ha"] for x in s], color="#b3401a", lw=2, label="Burned area (ha)")
    ax.plot(t, [x["burning_ha"] for x in s], color="#e8a317", lw=1.6, ls="--", label="Burning area (ha)")
    ax.set_xlabel("Simulated time (min)")
    ax.set_ylabel("ha")
    ax.grid(alpha=.3)
    ax.legend(frameon=False, fontsize=8)
    ax.set_title("SIMULATED burned and burning area", fontsize=9)
    fig.tight_layout()
    b = io.BytesIO()
    fig.savefig(b, format="png")
    plt.close(fig)
    out["area"] = b.getvalue()
    fig, ax = plt.subplots(figsize=(6.4, 1.9), dpi=150)
    ax.plot(t, [x["front_m"] for x in s], color="#2f6db3", lw=2, label="Front distance (m)")
    ax2 = ax.twinx()
    ax2.plot(t, [x["ros"] for x in s], color="#7a7a7a", lw=1.2, label="Rate of spread (m/min)")
    ax.set_xlabel("Simulated time (min)")
    ax.set_ylabel("m")
    ax2.set_ylabel("m/min")
    ax.grid(alpha=.3)
    ax.set_title("SIMULATED front distance and rate of spread", fontsize=9)
    fig.tight_layout()
    b = io.BytesIO()
    fig.savefig(b, format="png")
    plt.close(fig)
    out["front"] = b.getvalue()
    return out


# ── content model shared by PDF / DOCX / HTML ───────────────────────────────

def _ist(v) -> str:
    if not v:
        return "not available"
    from src.utils.timezone import format_ist
    return format_ist(v)


def sections(snap: dict) -> List[dict]:
    """[{title, paras:[str], tables:[(header, rows)], charts:[key]}] - one per page."""
    loc, w, r, f, res = snap["location"], snap["weather"], snap["risk"], snap["firms"], snap["results"]
    mode = snap["mode"]
    interp = (f"In this {'hypothetical WHAT-IF scenario' if mode != 'LIVE' else 'LIVE simulation started from satellite '}"
              f"{'' if mode != 'LIVE' else 'fire detections'}, the simulated fire burned {res['burned_ha']:.1f} ha in "
              f"{res['duration_label']}, its front reaching {res['front_distance_m'] or 0:.0f} m from the ignition "
              f"and spreading towards the {res['spread_towards']} with the wind. "
              + (f"It went out by itself at T+{res['extinguished_at_min']:.0f} min. " if res.get("extinguished_at_min")
                 is not None else "It was still active at the end of the simulated time. ")
              + ("The fire left the focus area. " if res.get("left_focus") else ""))
    risk_txt = (f"{r.get('category') or 'category not available'} - {r['score_pct']}% (PREDICTED by "
                f"{r.get('model') or 'the ML model'}, zone {r.get('zone')})" if r.get("score_pct") is not None
                else "not available")
    p1 = {"title": "1. Executive summary", "paras": [interp, DISCLAIMER], "tables": [(
        ("Item", "Value", "Type"),
        [("Report ID", snap["report_id"], "-"), ("Generated", _ist(snap["generated_utc"]), "-"),
         ("Forest / region", f"{loc['name']} / {snap.get('region') or '-'}", "-"),
         ("Coordinates", f"{loc['lat']:.5f} N, {loc['lon']:.5f} E", "-"),
         ("Focus area", f"{snap['focus']['width_m']:.0f} × {snap['focus']['height_m']:.0f} m "
                        f"({snap['focus']['cell_m']:.0f} m cells)", "-"),
         ("Simulation domain", f"{snap['domain_initial']['width_m'] / 1000:.2f} × {snap['domain_initial']['height_m'] / 1000:.2f}"
                               f" km initial → {snap['domain_final']['width_m'] / 1000:.2f} × "
                               f"{snap['domain_final']['height_m'] / 1000:.2f} km final", "SIMULATED"),
         ("Mode", "LIVE" if mode == "LIVE" else f"{mode} (hypothetical)", "-"),
         ("Ignition", snap["ignition"]["kind"] + (f", {len(snap['ignition']['points'])} point(s)"
                                                  if snap["ignition"]["points"] else ""),
          "OBSERVED" if mode == "LIVE" else "SCENARIO INPUT"),
         ("Temperature / humidity", f"{w.get('temp_c', '-')} °C / {w.get('humidity_pct', '-')} %", w["label"].split(" ")[0]),
         ("Wind used", f"{snap['wind_used']['speed_ms']} m/s from {compass(snap['wind_used']['from_deg'])} "
                       f"({snap['wind_used']['from_deg']:.0f}°)", w["label"].split(" ")[0]),
         ("ML fire risk", risk_txt, "PREDICTED")])], "charts": []}
    det_rows = [(f"{d.get('lat', 0):.4f}, {d.get('lon', 0):.4f}", str(d.get("acq_utc") or d.get("date") or "-"),
                 str(d.get("confidence") or d.get("conf") or "-"), str(d.get("satellite") or d.get("sat") or
                                                                    d.get("instrument") or "-"))
                for d in f.get("detections") or []]
    hist = snap["historical"]
    fwi = snap["fwi"]
    p2 = {"title": "2. Data and predictive analysis", "paras": [
        f"Weather ({w['label']}): source {w.get('source')}, time {_ist(w.get('timestamp_utc'))}, status "
        f"{w.get('status')}.",
        (f"NASA FIRMS: status {f.get('status')}, fetched {_ist(f.get('fetched_utc'))}, latest acquisition "
         f"{_ist(f.get('latest_acq_utc'))}; {f.get('n_total', 0)} detection(s) in the search box, "
         f"{f.get('n_in_area', 0)} in the simulation area, {f.get('n_valid', 0)} used as ignitions. Hotspots are "
         "satellite OBSERVATIONS, not ground-verified fires." if f.get("status") not in (None,) and
         not str(f.get("status")).startswith("not used") else
         "NASA FIRMS: not part of this run (WHAT-IF hypothetical ignition) - no observed fire is claimed."),
        (f"Historical fire records: {hist['note']} Zone fire rate {hist.get('zone_fire_rate')}, neighbourhood "
         f"{hist.get('neighbour_fire_rate')} (source: {hist.get('source')})." if hist.get("available")
         else "Historical fire records: " + hist["note"]),
        f"ML model: {r.get('model') or 'model details not available in this snapshot'}"
        + (f" ({r['model_file']})" if r.get("model_file") else "") + f"; prediction for grid zone {r.get('zone')}"
        + (f" ({r['zone_distance_km']} km from the location)" if r.get("zone_distance_km") is not None else "")
        + f": {risk_txt}."],
        "tables": [(("Canadian FWI component", "Value", "Type"),
                    [(k.upper(), "not calculated" if v is None else f"{v}", "DERIVED") for k, v in fwi.items()])]
        + ([(("Detection (lat, lon)", "Acquired (UTC)", "Confidence", "Sensor"), det_rows)] if det_rows else []),
        "charts": []}
    mdl = snap["model"]
    p3 = {"title": "3. Simulation results (SIMULATED)", "paras": [
        f"Fire model: {mdl.get('model') or '-'}; fuel {mdl.get('fuel_model') or '-'}; head / flank / back rate of "
        f"spread {mdl.get('ros_head_m_per_min')} / {mdl.get('ros_flank_m_per_min')} / {mdl.get('ros_back_m_per_min')} "
        f"m/min; fire-ellipse length:breadth {mdl.get('length_to_breadth')}; terrain {snap.get('terrain')}.",
        "Operational interpretation: use the simulated extent and direction to discuss where a fire starting at "
        "the selected point could reach under these conditions; confirm any real fire through official channels."],
        "tables": [(("Result", "Value"),
                    [("Simulated duration", res["duration_label"]), ("Total burned area", f"{res['burned_ha']:.2f} ha"),
                     ("Active burning area (end)", f"{res['burning_ha']:.2f} ha"),
                     ("Fire perimeter", f"{res['perimeter_m'] or 0:.0f} m"),
                     ("Front distance", f"{res['front_distance_m'] or 0:.0f} m"),
                     ("Peak rate of spread", f"{res['peak_ros']:.1f} m/min"),
                     ("Max fireline intensity", f"{res['max_kw_m']:,.0f} kW/m" if res.get("max_kw_m") else "-"),
                     ("Wind / spread direction", f"from {compass(snap['wind_used']['from_deg'])} → towards "
                                                 f"{res['spread_towards']}"),
                     ("Domain expansions", str(snap["domain_final"]["expansions"])),
                     ("Fire state at end", (f"out at T+{res['extinguished_at_min']:.0f} min"
                                            if res.get("extinguished_at_min") is not None else "still active"))]),
                   (("Model limitation",), [(x,) for x in snap["limitations"]])],
        "charts": ["area", "front"]}
    return [p1, p2, p3]


# ── PDF ─────────────────────────────────────────────────────────────────────

def render_pdf(snap: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=15, spaceAfter=4,
                        textColor=colors.HexColor("#7a2a0c"))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=4,
                        spaceAfter=4, textColor=colors.HexColor("#22313f"))
    ch = charts(snap)
    buf = io.BytesIO()
    banner = ("SIMULATION ONLY - hypothetical scenario, not a confirmed fire" if snap["mode"] != "LIVE"
              else "LIVE - started from NASA FIRMS satellite detections; spread is SIMULATED")

    def deco(c, d):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(15 * mm, 10 * mm, f"{TITLE} · Report {snap['report_id']} · {banner}")
        c.drawRightString(195 * mm, 10 * mm, f"Page {d.page}")
        c.restoreState()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm,
                            bottomMargin=17 * mm, title=f"Simulation report {snap['report_id']}", author=TITLE)
    story = [Paragraph(escape(TITLE), h1), Paragraph(f"Fire-spread simulation report · <b>{escape(snap['report_id'])}</b>"
                                                     f" · {escape(_ist(snap['generated_utc']))}", body),
             Paragraph(f"<font color='#b3401a'><b>{escape(banner)}</b></font>", body), Spacer(1, 6)]
    for k, sec in enumerate(sections(snap)):
        if k:
            story.append(PageBreak())
        story.append(Paragraph(escape(sec["title"]), h2))
        for para in sec["paras"]:
            story += [Paragraph(escape(para), body), Spacer(1, 3)]
        for header, rows in sec["tables"]:
            data = [[Paragraph(f"<b>{escape(str(x))}</b>", small) for x in header]] + \
                   [[Paragraph(escape(str(c)), small) for c in row] for row in rows]
            widths = {1: [180 * mm], 2: [55 * mm, 125 * mm], 3: [48 * mm, 100 * mm, 32 * mm],
                      4: [45 * mm, 55 * mm, 35 * mm, 45 * mm]}[len(header)]
            t = Table(data, colWidths=widths, repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1e4dc")),
                                   ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b9b9b9")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#faf7f5")])]))
            story += [t, Spacer(1, 6)]
        for key in sec["charts"]:
            if key in ch:
                story += [Image(io.BytesIO(ch[key]), width=160 * mm, height=160 * mm * (2.1 if key == "area" else 1.9) / 6.4),
                          Spacer(1, 4)]
        if sec["charts"] and not ch:
            story.append(Paragraph("Charts unavailable (matplotlib not installed or fewer than 2 recorded frames).", small))
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


# ── DOCX ────────────────────────────────────────────────────────────────────

def render_docx(snap: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401  (portrait default)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(15)
        s.top_margin, s.bottom_margin = Mm(14), Mm(16)
        # page number field in the footer
        p = s.footer.paragraphs[0]
        p.text = f"{TITLE} · Report {snap['report_id']} · page "
        r = p.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.text = txt
            r._r.append(el)
        for run in p.runs:
            run.font.size = Pt(7.5)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(9.5)
    h = d.add_heading(TITLE, level=0)
    h.runs[0].font.size = Pt(18)
    d.add_paragraph(f"Fire-spread simulation report · {snap['report_id']} · {_ist(snap['generated_utc'])}")
    ban = d.add_paragraph()
    run = ban.add_run("SIMULATION ONLY - hypothetical scenario, not a confirmed fire" if snap["mode"] != "LIVE"
                      else "LIVE - started from NASA FIRMS satellite detections; spread is SIMULATED")
    run.bold, run.font.color.rgb = True, RGBColor(0xB3, 0x40, 0x1A)
    ch = charts(snap)
    for k, sec in enumerate(sections(snap)):
        if k:
            d.add_page_break()
        d.add_heading(sec["title"], level=1)
        for para in sec["paras"]:
            d.add_paragraph(para)
        for header, rows in sec["tables"]:
            t = d.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 2"
            for i, x in enumerate(header):
                t.rows[0].cells[i].text = str(x)
            for row in rows:
                cells = t.add_row().cells
                for i, x in enumerate(row):
                    cells[i].text = str(x)
            d.add_paragraph()
        for key in sec["charts"]:
            if key in ch:
                d.add_picture(io.BytesIO(ch[key]), width=Mm(170))
        if sec["charts"] and not ch:
            d.add_paragraph("Charts unavailable (matplotlib not installed or fewer than 2 recorded frames).")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


# ── printable HTML ──────────────────────────────────────────────────────────

def render_html(snap: dict, auto_print: bool = False) -> str:
    import base64
    from html import escape
    ch = charts(snap)
    parts = [f"<h1>{escape(TITLE)}</h1><p>Fire-spread simulation report · <b>{escape(snap['report_id'])}</b> · "
             f"{escape(_ist(snap['generated_utc']))}</p>"]
    for k, sec in enumerate(sections(snap)):
        parts.append(f"<section class='{'pb' if k else ''}'><h2>{escape(sec['title'])}</h2>")
        parts += [f"<p>{escape(p)}</p>" for p in sec["paras"]]
        for header, rows in sec["tables"]:
            parts.append("<table><tr>" + "".join(f"<th>{escape(str(x))}</th>" for x in header) + "</tr>" +
                         "".join("<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in row) + "</tr>" for row in rows)
                         + "</table>")
        for key in sec["charts"]:
            if key in ch:
                parts.append(f"<img src='data:image/png;base64,{base64.b64encode(ch[key]).decode()}'>")
        parts.append("</section>")
    css = ("body{font-family:Segoe UI,Arial,sans-serif;font-size:12px;color:#222;background:#fff;margin:14px}"
           "h1{font-size:19px;color:#7a2a0c;margin:0 0 4px}h2{font-size:15px;border-bottom:1px solid #ccc}"
           "table{border-collapse:collapse;width:100%;margin:6px 0}td,th{border:1px solid #bbb;padding:3px 5px;"
           "vertical-align:top;text-align:left}th{background:#f1e4dc}img{width:100%;max-width:720px}"
           "@media print{.pb{page-break-before:always}button{display:none}}")
    btn = "<button onclick='window.print()' style='padding:6px 12px;margin:6px 0'>Print report</button>"
    js = "<script>setTimeout(()=>window.print(),400)</script>" if auto_print else ""
    return f"<html><head><meta charset='utf-8'><style>{css}</style></head><body>{btn}{''.join(parts)}{js}</body></html>"
